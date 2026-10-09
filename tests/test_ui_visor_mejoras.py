"""Regresiones de búsqueda, enfoque de tabla y separación de CFDI por visor."""
from pathlib import Path
from types import SimpleNamespace
import tkinter as tk

import customtkinter as ctk
import pytest

from conxml import estado_local
from conxml.catalog.db import Catalogo
from conxml.catalog.importer import importar_carpeta
from conxml.ui.app import ConXmlApp


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setenv('CONXML_DATA_DIR', str(tmp_path))
    monkeypatch.setattr('conxml.ui.app.Actualizaciones', lambda app: None)
    with Catalogo(tmp_path / 'catalogo.db') as cat:
        cat.crear_cliente('C1', 'Cliente de prueba', 'EKU9003173C9')
        importar_carpeta(cat, Path(__file__).parent / 'fixtures', 'C1')
    estado_local.guardar({'sesion': {'cliente': 'C1', 'pantallas': {
        'cfdi40': {'filtros': dict.fromkeys(('uuid', 'rfc', 'serie', 'folio'), '')}
    }}})
    try:
        root = ctk.CTk()
    except tk.TclError as exc:
        pytest.skip(f'Sin display: {exc}')
    errores = []
    root.report_callback_exception = lambda *args: errores.append(args)
    ventana = ConXmlApp(root, cliente_actual='C1')
    monkeypatch.setattr(ventana, 'ejecutar', lambda *args, **kwargs: None)
    root.geometry('1400x900')
    root.update()
    yield ventana
    root.destroy()
    assert not errores, errores


def test_busqueda_vacia_conserva_ayuda_y_fecha_restaurable(app):
    app.navegar('admin40')
    pantalla = app._pantallas['admin40']
    app.master.update()
    for clave, entrada in pantalla._filtros.items():
        assert entrada.get() == ''
        assert entrada._entry.get() == entrada.cget('placeholder_text')
    assert pantalla._fecha.get() == ''
    assert pantalla._fecha._entrada._entry.get() == 'DD/MM/AA'
    assert [w.cget('text') for w in pantalla._etiquetas_filtro] == [
        'UUID', 'RFC', 'Serie', 'Folio', 'Fecha · DD/MM/AA', 'Monto']
    pantalla._fecha.set('07/10/26')
    app.guardar_sesion()
    app.recargar_datos_locales()
    nueva = app.master._conxml_app
    assert nueva._pantallas['admin40']._fecha.get() == '07/10/26'


def test_ampliar_conserva_filas_seleccion_y_restaura_al_salir(app):
    app.navegar('admin40')
    p = app._pantallas['admin40']
    app.master.update()
    filas = p._tabla.get_children()
    p._tabla.selection_set(filas[0])
    tamano = (p._tabla.winfo_width(), p._tabla.winfo_height())
    p._alternar_tabla_ampliada()
    app.master.update()
    assert p._tabla.winfo_width() > tamano[0]
    assert p._tabla.winfo_height() > tamano[1]
    assert p._tabla.selection() == (filas[0],)
    assert p._tabla.get_children() == filas
    assert not app.sidebar_visible
    app._restaurar_tabla()
    app.master.update()
    assert app.sidebar_visible
    assert not p._tabla_ampliada
    p._alternar_tabla_ampliada()
    app.navegar('nomina')
    assert not p._tabla_ampliada and app.sidebar_visible


def test_cada_visor_consulta_solo_su_tipo_y_totales(app):
    for clave, tipos in [('admin40', {'I', 'E', 'T'}), ('nomina', {'N'}), ('pagos', {'P'})]:
        app.navegar(clave)
        pantalla = app._pantallas[clave]
        with Catalogo(app.db_path) as cat:
            filas = list(pantalla._consulta(cat))
            esperadas = {f['uuid'] for f in cat.consulta(cliente='C1') if f['tipo_comprobante'] in tipos}
        assert {f['uuid'] for f in filas} == esperadas
        assert pantalla._uuids_visibles() <= esperadas
        if clave != 'pagos':
            assert {f['uuid'] for f in pantalla._totales_panel._filas} == esperadas


def test_doble_clic_abre_pdf_de_la_fila_y_no_del_encabezado(app, monkeypatch):
    app.navegar('admin40')
    p = app._pantallas['admin40']
    app.master.update()
    fila = p._tabla.get_children()[1]
    p._tabla.see(fila)
    x, y, ancho, alto = p._tabla.bbox(fila)
    abiertas = []
    monkeypatch.setattr(p, '_vista_previa_pdf', lambda: abiertas.extend(p._documentos_pdf(individual=True)))
    p._doble_clic_pdf(SimpleNamespace(x=x + 10, y=y + alto // 2))
    assert len(abiertas) == 1
    assert abiertas[0]['uuid'] in p._uuids_visibles(seleccion=True)
    p._doble_clic_pdf(SimpleNamespace(x=10, y=2))
    assert len(abiertas) == 1
    assert all(b.cget('text') not in ('Vista previa PDF', 'Guardar PDF') for b in p.botones)


def test_validacion_usa_seleccion_o_vista_filtrada(app, monkeypatch):
    app.navegar('admin40')
    p = app._pantallas['admin40']
    filas = p._tabla.get_children()
    p._tabla.selection_set(filas[0])
    llamadas = []
    monkeypatch.setattr(p, '_run_validar', lambda cliente, force, config, progreso, uuids: llamadas.append((config, uuids)))
    monkeypatch.setattr(app, 'ejecutar', lambda fn, cb, texto, **kw: fn(lambda *_: None))
    p._validar()
    assert llamadas[0][0].trabajadores == 3
    assert llamadas[0][1] == p._uuids_visibles(seleccion=True)
    assert len(llamadas[0][1]) == 1
    p._tabla.selection_remove(*p._tabla.selection())
    p._validar()
    assert llamadas[1][1] == p._uuids_visibles()


def test_exportar_xml40_excluye_nomina_y_pagos(app, tmp_path):
    from openpyxl import load_workbook
    p = app._pantallas['admin40']
    destino = p._run_exportar(tmp_path / 'xml40.xlsx', 'C1')
    wb = load_workbook(destino, read_only=True)
    filas = list(wb['Listado'].iter_rows(min_row=2, values_only=True))
    with Catalogo(app.db_path) as cat:
        esperados = {f['uuid'] for f in cat.consulta(cliente='C1', tipos=('I', 'E', 'T'))}
    assert {fila[5] for fila in filas} == esperados
    wb.close()


def test_busqueda_pagos_no_muestra_rep_que_no_coincide(app):
    app.navegar('pagos')
    p = app._pantallas['pagos']
    p._filtros['rfc'].insert(0, 'NO-EXISTE')
    for vista in ('conciliacion', 'facturas_p', 'pagos', 'doctos'):
        p._vista = vista
        p._configurar_columnas(p._columnas_actuales())
        p._cargar_tabla()
        assert not p._tabla.get_children()
